"""
ATS Resume Checker
------------------
Upload a resume (PDF / DOCX / TXT) and get:
  * an estimated ATS score (0-100)
  * category scores, strengths and weaknesses
  * missing keywords, prioritised improvements and bullet rewrites
  * optional match against a pasted job description

UI: Streamlit  |  AI: Google Gemini Flash (google-genai SDK)
"""

import io
import json
import os
import re
import time

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
DEFAULT_MODEL = "gemini-3.8-flash"
FALLBACK_MODELS = ["gemini-3.7-flash"]  # tried if the main model is overloaded or unavailable
RETRY_DELAYS = [2, 5]  # seconds to wait between retries on the same model
MAX_RESUME_CHARS = 15000  # keeps the prompt small and fast
MAX_JD_CHARS = 6000

# Weights used to compute the overall ATS score from the category scores.
CATEGORY_WEIGHTS = {
    "keywords_relevance": 0.30,
    "formatting_structure": 0.20,
    "content_impact": 0.25,
    "readability_clarity": 0.10,
    "completeness": 0.15,
}
CATEGORY_LABELS = {
    "keywords_relevance": "Keywords & relevance",
    "formatting_structure": "Formatting & structure",
    "content_impact": "Content & impact",
    "readability_clarity": "Readability & clarity",
    "completeness": "Completeness",
}


# --------------------------------------------------------------------------- #
# Text extraction
# --------------------------------------------------------------------------- #
def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT file."""
    name = filename.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:  # resumes often use tables for layout
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# --------------------------------------------------------------------------- #
# Quick rule-based checks (no AI needed)
# --------------------------------------------------------------------------- #
def quick_checks(text: str) -> list[dict]:
    """Cheap, deterministic checks that ATS parsers care about."""
    lower = text.lower()
    words = re.findall(r"\b\w+\b", text)
    word_count = len(words)

    has_email = bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text))
    has_phone = any(
        10 <= len(re.sub(r"\D", "", m)) <= 15
        for m in re.findall(r"\+?\d[\d\s().-]{8,}\d", text)
    )
    has_linkedin = "linkedin.com" in lower
    bullets = len(re.findall(r"(?m)^\s*[•\-\u2022\u25AA\u25CF\*]\s+", text))
    numbers = len(re.findall(r"\d+(?:[.,]\d+)?\s*(?:%|\+|k\b|m\b|x\b)|[$€£]\s*\d", lower))

    def has_section(*names: str) -> bool:
        return any(re.search(rf"(?m)^\s*(?:professional |work |relevant |technical )?{n}\b", lower) for n in names)

    return [
        {"check": "Email address found", "ok": has_email},
        {"check": "Phone number found", "ok": has_phone},
        {"check": "LinkedIn / profile link found", "ok": has_linkedin},
        {
            "check": "Experience section heading",
            "ok": has_section("experience", "work experience", "employment", "work history"),
        },
        {"check": "Education section heading", "ok": has_section("education", "academic")},
        {"check": "Skills section heading", "ok": has_section("skills", "technical skills", "core competencies")},
        {"check": f"Length looks reasonable ({word_count} words; ideal 300-900)", "ok": 300 <= word_count <= 900},
        {"check": f"Uses bullet points ({bullets} found)", "ok": bullets >= 5},
        {"check": f"Quantified achievements ({numbers} found)", "ok": numbers >= 3},
    ]


# --------------------------------------------------------------------------- #
# Gemini call
# --------------------------------------------------------------------------- #
PROMPT_TEMPLATE = """You are an expert technical recruiter and ATS (Applicant Tracking System) specialist.
Evaluate the resume below as an ATS would, then give honest, specific, actionable feedback.

Rules:
- Base everything ONLY on the resume text (and the job description if one is provided). Never invent experience.
- Score each category from 0 to 100. Be strict and realistic; most resumes score between 45 and 85.
- "keywords_relevance": {kw_rule}
- Improvements must be specific (mention the section or line they apply to), not generic advice.
- "bullet_rewrites": pick up to 4 weak bullets from the resume and rewrite them with strong action verbs and
  measurable impact. If the real number is unknown, use a clearly marked placeholder like [X%].
- Return ONLY a valid JSON object. No markdown, no code fences, no commentary.

JSON schema:
{{
  "category_scores": {{
    "keywords_relevance": int,
    "formatting_structure": int,
    "content_impact": int,
    "readability_clarity": int,
    "completeness": int
  }},
  "summary": "2-3 sentence overall assessment",
  "strengths": ["..."],
  "weaknesses": ["..."],
  "missing_keywords": ["..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "...", "issue": "...", "suggestion": "..."}}
  ],
  "bullet_rewrites": [{{"original": "...", "improved": "..."}}],
  "jd_match_percent": {jd_field}
}}

Automatic checks already run on the resume (use them as hints):
{checks}

{jd_block}
RESUME TEXT:
\"\"\"
{resume}
\"\"\"
"""


def build_prompt(resume_text: str, job_description: str, checks: list[dict]) -> str:
    checks_str = "\n".join(f"- {c['check']}: {'PASS' if c['ok'] else 'FAIL'}" for c in checks)
    if job_description.strip():
        kw_rule = "how well the resume's keywords and skills match the provided job description."
        jd_field = "int (0-100, how well the resume matches the job description)"
        jd_block = f'JOB DESCRIPTION:\n"""\n{job_description.strip()[:MAX_JD_CHARS]}\n"""\n'
    else:
        kw_rule = "how well the resume uses relevant, industry-standard keywords and skills for the candidate's apparent target role."
        jd_field = "null"
        jd_block = "No job description was provided.\n"
    return PROMPT_TEMPLATE.format(
        kw_rule=kw_rule,
        jd_field=jd_field,
        checks=checks_str,
        jd_block=jd_block,
        resume=resume_text[:MAX_RESUME_CHARS],
    )


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating code fences or stray text."""
    if not raw:
        raise ValueError("The model returned an empty response.")
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip(), flags=re.IGNORECASE)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start != -1 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise ValueError("Could not parse the model's response as JSON.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def _str_list(value) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalize_result(data: dict) -> dict:
    """Validate and clean the model output; compute the overall score in code."""
    raw_scores = data.get("category_scores") or {}
    scores = {k: _clamp(raw_scores.get(k)) for k in CATEGORY_WEIGHTS}
    overall = round(sum(scores[k] * w for k, w in CATEGORY_WEIGHTS.items()))

    priority_order = {"high": 0, "medium": 1, "low": 2}
    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            priority = str(item.get("priority", "Medium")).strip().title()
            if priority.lower() not in priority_order:
                priority = "Medium"
            improvements.append(
                {
                    "priority": priority,
                    "section": str(item.get("section", "General")).strip(),
                    "issue": str(item.get("issue", "")).strip(),
                    "suggestion": str(item.get("suggestion", "")).strip(),
                }
            )
    improvements.sort(key=lambda i: priority_order[i["priority"].lower()])

    rewrites = []
    for item in data.get("bullet_rewrites") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append({"original": str(item["original"]), "improved": str(item["improved"])})

    jd = data.get("jd_match_percent")
    return {
        "overall_score": overall,
        "category_scores": scores,
        "summary": str(data.get("summary", "")).strip(),
        "strengths": _str_list(data.get("strengths")),
        "weaknesses": _str_list(data.get("weaknesses")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "improvements": improvements,
        "bullet_rewrites": rewrites,
        "jd_match_percent": None if jd is None else _clamp(jd),
    }


def _is_overloaded(exc: Exception) -> bool:
    msg = str(exc)
    return any(t in msg for t in ("503", "UNAVAILABLE", "high demand", "overloaded", "INTERNAL"))


def _is_model_missing(exc: Exception) -> bool:
    msg = str(exc)
    return "404" in msg or "NOT_FOUND" in msg


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str, checks: list[dict]) -> dict:
    """Call Gemini. Retries when overloaded (503) and falls back to another model if needed."""
    client = genai.Client(api_key=api_key)
    prompt = build_prompt(resume_text, job_description, checks)
    config = types.GenerateContentConfig(response_mime_type="application/json")

    models = [model] + [m for m in FALLBACK_MODELS if m != model]
    last_error: Exception = RuntimeError("No model was tried.")
    for name in models:
        for attempt in range(len(RETRY_DELAYS) + 1):
            try:
                response = client.models.generate_content(model=name, contents=prompt, config=config)
                return normalize_result(parse_json_response(response.text))
            except Exception as exc:
                last_error = exc
                if _is_model_missing(exc):
                    break  # this model name is gone: go straight to the next one
                if not _is_overloaded(exc):
                    raise  # e.g. invalid API key, bad request: retrying won't help
                if attempt < len(RETRY_DELAYS):
                    time.sleep(RETRY_DELAYS[attempt])
                # else: out of retries, move on to the next model
    raise last_error


# --------------------------------------------------------------------------- #
# Report export
# --------------------------------------------------------------------------- #
def build_markdown_report(result: dict, checks: list[dict]) -> str:
    lines = ["# ATS Resume Report\n", f"**Overall ATS score: {result['overall_score']}/100**\n"]
    if result["jd_match_percent"] is not None:
        lines.append(f"**Job description match: {result['jd_match_percent']}%**\n")
    lines += [result["summary"], "", "## Category scores"]
    lines += [f"- {CATEGORY_LABELS[k]}: {v}/100" for k, v in result["category_scores"].items()]
    lines += ["", "## Strengths"] + [f"- {s}" for s in result["strengths"]]
    lines += ["", "## Weaknesses"] + [f"- {s}" for s in result["weaknesses"]]
    lines += ["", "## Missing keywords", ", ".join(result["missing_keywords"]) or "None"]
    lines += ["", "## Improvements"]
    for i in result["improvements"]:
        lines.append(f"- **[{i['priority']}] {i['section']}** - {i['issue']} -> {i['suggestion']}")
    lines += ["", "## Bullet rewrites"]
    for r in result["bullet_rewrites"]:
        lines += [f"- Before: {r['original']}", f"  After: {r['improved']}"]
    lines += ["", "## Automatic checks"]
    lines += [f"- {'PASS' if c['ok'] else 'FAIL'}: {c['check']}" for c in checks]
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# UI helpers
# --------------------------------------------------------------------------- #
def get_api_key(sidebar_value: str) -> str:
    """Priority: sidebar input > Streamlit secrets > environment variable."""
    if sidebar_value.strip():
        return sidebar_value.strip()
    try:
        secret = st.secrets.get("GEMINI_API_KEY", "")
        if secret:
            return secret
    except Exception:  # no secrets file present
        pass
    return os.environ.get("GEMINI_API_KEY", "")


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Weak"


def render_results(result: dict, checks: list[dict]) -> None:
    st.divider()
    top_left, top_right = st.columns([1, 2])
    with top_left:
        st.metric("Estimated ATS score", f"{result['overall_score']} / 100")
        st.progress(result["overall_score"] / 100)
        st.caption(f"Rating: {score_label(result['overall_score'])}")
        if result["jd_match_percent"] is not None:
            st.metric("Job description match", f"{result['jd_match_percent']}%")
    with top_right:
        st.subheader("Summary")
        st.write(result["summary"] or "No summary returned.")

    st.subheader("Category scores")
    cols = st.columns(len(result["category_scores"]))
    for col, (key, value) in zip(cols, result["category_scores"].items()):
        col.metric(CATEGORY_LABELS[key], value)
        col.progress(value / 100)

    tab_fix, tab_kw, tab_rewrite, tab_sw, tab_checks = st.tabs(
        ["Improvements", "Missing keywords", "Bullet rewrites", "Strengths & weaknesses", "Automatic checks"]
    )

    with tab_fix:
        if not result["improvements"]:
            st.info("No improvements returned.")
        for item in result["improvements"]:
            icon = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}[item["priority"]]
            with st.expander(f"{icon} {item['priority']} - {item['section']}", expanded=item["priority"] == "High"):
                st.markdown(f"**Issue:** {item['issue']}")
                st.markdown(f"**Fix:** {item['suggestion']}")

    with tab_kw:
        if result["missing_keywords"]:
            st.write("Consider adding these (only where they truthfully apply to you):")
            st.write(" ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.success("No major missing keywords detected.")

    with tab_rewrite:
        if not result["bullet_rewrites"]:
            st.info("No rewrites returned.")
        for r in result["bullet_rewrites"]:
            st.markdown(f"**Before:** {r['original']}")
            st.markdown(f"**After:** {r['improved']}")
            st.divider()

    with tab_sw:
        left, right = st.columns(2)
        with left:
            st.markdown("#### Strengths")
            for s in result["strengths"]:
                st.markdown(f"- {s}")
        with right:
            st.markdown("#### Weaknesses")
            for s in result["weaknesses"]:
                st.markdown(f"- {s}")

    with tab_checks:
        for c in checks:
            st.markdown(f"{'✅' if c['ok'] else '❌'} {c['check']}")

    st.download_button(
        "Download report (.md)",
        data=build_markdown_report(result, checks),
        file_name="ats_resume_report.md",
        mime="text/markdown",
    )


# --------------------------------------------------------------------------- #
# Main app
# --------------------------------------------------------------------------- #
def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an estimated ATS score and concrete ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        key_input = st.text_input(
            "Gemini API key",
            type="password",
            help="Free key from https://aistudio.google.com/apikey. Not stored anywhere.",
        )
        model = st.text_input("Gemini model", value=DEFAULT_MODEL, help="Any Gemini Flash model name works.")
        st.markdown("---")
        st.caption(
            "The score is an AI-based estimate of ATS-friendliness, not the output of a real employer's ATS. "
            "Use it as guidance."
        )

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, for a tailored keyword match)",
        height=150,
        placeholder="Paste the job posting here...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        api_key = get_api_key(key_input)
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
        else:
            try:
                with st.spinner("Reading your resume..."):
                    text = extract_text(uploaded.name, uploaded.getvalue())
                if len(text.split()) < 30:
                    raise ValueError(
                        "Almost no text could be read. If your PDF is a scan or image, "
                        "export it as a text-based PDF or DOCX and try again. "
                        "(This is also exactly the problem real ATS systems have.)"
                    )
                checks = quick_checks(text)
                with st.spinner("Analyzing with Gemini..."):
                    result = analyze_resume(api_key, model.strip() or DEFAULT_MODEL, text, job_description, checks)
                st.session_state["result"] = result
                st.session_state["checks"] = checks
            except ValueError as exc:
                st.session_state.pop("result", None)
                st.error(str(exc))
            except Exception as exc:  # network, quota, invalid key, bad model name ...
                st.session_state.pop("result", None)
                if _is_overloaded(exc):
                    st.error(
                        "Google's Gemini servers are overloaded right now (temporary). "
                        "Please wait a minute and click Analyze again."
                    )
                else:
                    st.error(f"Analysis failed: {exc}")

    if "result" in st.session_state:
        render_results(st.session_state["result"], st.session_state["checks"])


if __name__ == "__main__":
    main()
